"""Export formatting (P6): columns, rows, CSV and XLSX bytes.

Pure functions: no database, no FastAPI, no storage. Exports are built from current
data on every request (R23) with the R24 columns, and every cell is escaped against
formula injection (engineering rule 6).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

__all__ = [
    "Column",
    "ColumnSource",
    "build_columns",
    "build_rows",
    "escape_cell",
    "export_filename",
    "to_csv",
    "to_xlsx",
]

#: Characters that make a spreadsheet treat a cell as a formula.
DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

ColumnSource = Literal["profile", "email", "answer", "meta"]

PROFILE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("name", "Name"),
    ("roll_no", "Roll No"),
    ("department", "Department"),
    ("year", "Year"),
    ("section", "Section"),
)
META_COLUMNS: tuple[tuple[str, str], ...] = (
    ("status", "Status"),
    ("registered_at", "Registered At"),
    ("attended", "Attended"),
    ("certificate_id", "Certificate ID"),
    ("certificate_status", "Certificate Status"),
)
_DEFAULT_WIDTHS = {
    "name": 28,
    "roll_no": 16,
    "department": 16,
    "year": 8,
    "section": 10,
    "email": 32,
    "status": 12,
    "registered_at": 27,
    "attended": 10,
    "certificate_id": 20,
    "certificate_status": 18,
}
_ANSWER_WIDTHS = {"long_text": 50, "multi_choice": 30, "short_text": 24, "email": 32}


@dataclass(frozen=True, slots=True)
class Column:
    """One Export column. `key` is the profile/meta name or the Form Field key."""

    key: str
    header: str
    source: ColumnSource
    field_type: str | None = None
    width: int = 16


def _get(obj: Any, name: str, default: Any = None) -> Any:
    if isinstance(obj, Mapping):
        return obj.get(name, default)
    return getattr(obj, name, default)


def escape_cell(value: Any) -> Any:
    """Neutralise formula injection: prefix a dangerous str with `'`. None becomes "".

    Non-str values are returned unchanged. Idempotent (an escaped value starts with `'`).
    """
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith(DANGEROUS_PREFIXES):
        return "'" + value
    return value


def build_columns(form_fields: Iterable[Any]) -> list[Column]:
    """Return the R24 columns in order.

    Profile Snapshot (name, roll_no, department, year, section), email, one column per
    Form Field in `position` order (hidden fields included, header suffixed
    " (hidden)"), then status, registered_at, attended, certificate_id,
    certificate_status. Fields without a `position` keep their given order.
    """
    columns = [
        Column(key, header, "profile", width=_DEFAULT_WIDTHS[key])
        for key, header in PROFILE_COLUMNS
    ]
    columns.append(Column("email", "Email", "email", width=_DEFAULT_WIDTHS["email"]))

    indexed = list(enumerate(form_fields))
    indexed.sort(
        key=lambda pair: (
            _get(pair[1], "position") if _get(pair[1], "position") is not None else pair[0],
            pair[0],
        )
    )
    for _, field in indexed:
        label = _get(field, "label") or _get(field, "key")
        header = f"{label} (hidden)" if _get(field, "hidden", False) else str(label)
        ftype = _get(field, "type")
        width = max(_ANSWER_WIDTHS.get(ftype, 16), min(len(header) + 2, 50))
        columns.append(Column(_get(field, "key"), header, "answer", ftype, width))

    columns.extend(
        Column(key, header, "meta", width=_DEFAULT_WIDTHS[key]) for key, header in META_COLUMNS
    )
    return columns


def _fmt_datetime(value: datetime) -> str:
    dt = value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
    return dt.isoformat()


def _yes_no(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def _format_value(value: Any, field_type: str | None = None) -> Any:
    if value is None:
        return ""
    if field_type == "yes_no" or isinstance(value, bool):
        return _yes_no(value)
    if isinstance(value, list | tuple):
        return "; ".join("" if v is None else str(v) for v in value)
    if isinstance(value, datetime):
        return _fmt_datetime(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Mapping):
        return str(dict(value))
    return value


def build_rows(
    registrations: Iterable[Mapping[str, Any]], columns: Sequence[Column]
) -> Iterator[list]:
    """Yield one escaped row per Registration, matching `columns`.

    Each Registration mapping has `profile_snapshot`, `email`, `answers`, `status`,
    `submitted_at`, `attended` (bool | None), `certificate_id`, `certificate_status`.
    """
    for reg in registrations:
        snapshot = reg.get("profile_snapshot") or {}
        answers = reg.get("answers") or {}
        row: list[Any] = []
        for col in columns:
            if col.source == "profile":
                value = _format_value(snapshot.get(col.key))
            elif col.source == "email":
                value = _format_value(reg.get("email"))
            elif col.source == "answer":
                value = _format_value(answers.get(col.key), col.field_type)
            elif col.key == "registered_at":
                value = _format_value(reg.get("submitted_at"))
            elif col.key == "attended":
                value = _yes_no(reg.get("attended"))
            else:
                value = _format_value(reg.get(col.key))
            row.append(escape_cell(value))
        yield row


def to_csv(columns: Sequence[Column], rows: Iterable[Sequence[Any]]) -> bytes:
    """CSV bytes, UTF-8 with a BOM so Excel shows non-ASCII names correctly."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([escape_cell(c.header) for c in columns])
    for row in rows:
        writer.writerow([escape_cell(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")


def _xlsx_value(value: Any) -> Any:
    value = escape_cell(value)
    if isinstance(value, str):
        # openpyxl refuses control characters that XML can't hold.
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def to_xlsx(columns: Sequence[Column], rows: Iterable[Sequence[Any]]) -> bytes:
    """XLSX bytes (write-only, streams rows): bold frozen header, per-column widths."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Registrations")
    for index, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = col.width
    ws.freeze_panes = "A2"

    bold = Font(bold=True)
    header = []
    for col in columns:
        cell = WriteOnlyCell(ws, value=_xlsx_value(col.header))
        cell.font = bold
        header.append(cell)
    ws.append(header)
    for row in rows:
        ws.append([_xlsx_value(v) for v in row])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_SLUG_STRIP_RE = re.compile(r"[^a-z0-9]+")


def export_filename(event_name: str, fmt: str, now: datetime | None = None) -> str:
    """A safe filename like `tech-fest-2026_registrations_2026-10-25.xlsx` (UTC date)."""
    fmt = fmt.lower().lstrip(".")
    if fmt not in ("csv", "xlsx"):
        raise ValueError(f"Unsupported export format: {fmt!r}")
    ascii_name = unicodedata.normalize("NFKD", event_name or "").encode("ascii", "ignore").decode()
    slug = _SLUG_STRIP_RE.sub("-", ascii_name.lower()).strip("-")[:60].strip("-") or "event"
    moment = now if now is not None else datetime.now(UTC)
    moment = moment.replace(tzinfo=UTC) if moment.tzinfo is None else moment.astimezone(UTC)
    return f"{slug}_registrations_{moment.date().isoformat()}.{fmt}"
