"""Drive the whole product end to end over HTTP, like the frontend would.

Needs a running API server and worker that share the same Postgres, Redis and outbox:

    uv run python scripts/e2e.py --base-url http://localhost:8002 --outbox <outbox dir>

Everything goes through the API. There are only two exceptions: the Platform Admin is
made with `scripts/create_admin.py`, and emails are read from the outbox folder. Each run
uses a unique suffix, so it can run again and again on the same database. It prints one
numbered PASS/FAIL line per check and exits non-zero if any check failed.
"""

import argparse
import contextlib
import csv
import email
import io
import secrets
import sys
import threading
import time
from collections import Counter
from datetime import UTC, datetime, timedelta
from email import policy
from pathlib import Path
from typing import Any

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parent))

from create_admin import create_admin  # noqa: E402

PASSWORD = "Passw0rd-e2e"
NEW_PASSWORD = "N3wPassw0rd-e2e"
ADMIN_EMAIL = "e2e-admin@uni-e2e.edu"
ADMIN_PASSWORD = "Adm1nPassw0rd-e2e"
DOMAIN = "uni-e2e.edu"
VERIFY_FIELDS = {
    "public_id",
    "name",
    "event_name",
    "club_name",
    "event_date",
    "status",
    "superseded_by",
    "revoked_reason",
    "revoked_at",
}


# --------------------------------------------------------------------------- reporting


class Report:
    def __init__(self) -> None:
        self.results: list[tuple[int, bool, str, str]] = []
        self.lock = threading.Lock()

    def check(self, name: str, ok: bool, detail: str = "") -> bool:
        with self.lock:
            n = len(self.results) + 1
            self.results.append((n, bool(ok), name, detail))
        line = f"{n:3d}. {'PASS' if ok else 'FAIL'}  {name}"
        if not ok and detail:
            line += f"\n       -> {detail}"
        print(line, flush=True)
        return bool(ok)

    def status(self, name: str, resp: httpx.Response, expected: int | tuple[int, ...]) -> bool:
        want = expected if isinstance(expected, tuple) else (expected,)
        return self.check(name, resp.status_code in want, describe(resp, want))

    def summary(self) -> int:
        passed = sum(1 for r in self.results if r[1])
        failed = len(self.results) - passed
        print("\n" + "=" * 70)
        print(f"Summary: {passed} passed, {failed} failed, {len(self.results)} checks")
        if failed:
            print("Failed checks:")
            for n, ok, name, detail in self.results:
                if not ok:
                    print(f"  {n}. {name}: {detail}")
        return 1 if failed else 0


def describe(resp: httpx.Response, want: tuple[int, ...] | None = None) -> str:
    body = resp.text
    if len(body) > 400:
        body = body[:400] + "..."
    want_text = f" (expected {'/'.join(map(str, want))})" if want else ""
    return f"{resp.request.method} {resp.request.url.path} -> {resp.status_code}{want_text}: {body}"


class StepFailed(Exception):
    """A step can't go on; the rest of the step is skipped."""


# --------------------------------------------------------------------------- actors


class Actor:
    """One person with their own cookie jar."""

    def __init__(self, api: str, email_addr: str, profile: dict[str, Any]) -> None:
        self.api = api
        self.email = email_addr
        self.profile = profile
        self.client = httpx.Client(base_url=api, timeout=60)
        self.id: str | None = None

    def signup(self) -> httpx.Response:
        resp = self.client.post(
            "/auth/signup", json={**self.profile, "email": self.email, "password": PASSWORD}
        )
        if resp.status_code == 201:
            self.id = resp.json()["id"]
        return resp

    def login(self, password: str) -> httpx.Response:
        return self.client.post("/auth/login", json={"email": self.email, "password": password})

    def get(self, path: str, **kw: Any) -> httpx.Response:
        return self.client.get(path, **kw)

    def post(self, path: str, **kw: Any) -> httpx.Response:
        return self.client.post(path, **kw)

    def put(self, path: str, **kw: Any) -> httpx.Response:
        return self.client.put(path, **kw)

    def delete(self, path: str, **kw: Any) -> httpx.Response:
        return self.client.delete(path, **kw)

    def close(self) -> None:
        self.client.close()


# --------------------------------------------------------------------------- outbox


def read_outbox(outbox: Path) -> list[email.message.EmailMessage]:
    if not outbox.exists():
        return []
    messages = []
    for path in sorted(outbox.glob("*.eml")):
        try:
            messages.append(email.message_from_bytes(path.read_bytes(), policy=policy.default))
        except Exception:  # noqa: BLE001 - a half-written file; read it next poll
            continue
    return messages


def body_text(msg: email.message.EmailMessage) -> str:
    part = msg.get_body(preferencelist=("plain", "html"))
    return part.get_content() if part is not None else ""


def has_attachment(msg: email.message.EmailMessage) -> bool:
    return any(True for _ in msg.iter_attachments())


# --------------------------------------------------------------------------- helpers


def make_poster_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (2400, 3600), (30, 60, 120))
    draw = ImageDraw.Draw(img)
    draw.rectangle((100, 100, 2300, 600), fill=(240, 200, 40))
    draw.text((200, 300), "ORGniser e2e poster", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_template_pdf(tags: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    width, height = landscape(A4)
    c = canvas.Canvas(buf, pagesize=(width, height))
    c.setFont("Helvetica-Bold", 28)
    c.drawCentredString(width / 2, height - 80, "Certificate of Participation")
    c.setFont("Helvetica", 18)
    y = height - 160
    for tag in tags:
        c.drawCentredString(width / 2, y, "{{" + tag + "}}")
        y -= 50
    c.showPage()
    c.save()
    return buf.getvalue()


def decode_qr(png: bytes, area: dict[str, int]) -> str:
    import cv2
    import numpy as np

    img = cv2.imdecode(np.frombuffer(png, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return ""
    detector = cv2.QRCodeDetector()
    text, _, _ = detector.detectAndDecode(img)
    if text:
        return text
    m = 150
    x0, y0 = max(area["x"] - m, 0), max(area["y"] - m, 0)
    x1 = min(area["x"] + area["w"] + m, img.shape[1])
    y1 = min(area["y"] + area["h"] + m, img.shape[0])
    text, _, _ = detector.detectAndDecode(img[y0:y1, x0:x1])
    return text or ""


def poll(fn, done, timeout: float, interval: float = 1.0):
    """Call `fn()` until `done(result)` or the timeout; return the last result."""
    deadline = time.monotonic() + timeout
    result = fn()
    while not done(result) and time.monotonic() < deadline:
        time.sleep(interval)
        result = fn()
    return result


def csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def xlsx_rows(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- the run


class E2E:
    def __init__(self, base_url: str, outbox: Path, frontend_base_url: str) -> None:
        self.api = base_url.rstrip("/") + "/api/v1"
        self.outbox = outbox
        self.frontend = frontend_base_url.rstrip("/")
        self.r = Report()
        self.suffix = datetime.now(UTC).strftime("%m%d%H%M%S") + secrets.token_hex(2)
        self.actors: list[Actor] = []
        self.s: dict[str, Any] = {}  # shared state between steps

    # -- small helpers

    def actor(self, label: str, dept: str, year: int, section: str) -> Actor:
        n = len(self.actors)
        a = Actor(
            self.api,
            f"{label}-{self.suffix}@{DOMAIN}",
            {
                "name": f"{label.title()} {self.suffix}",
                "roll_no": f"R{self.suffix}{n:03d}",
                "department": dept,
                "year": year,
                "section": section,
            },
        )
        self.actors.append(a)
        return a

    def need(self, ok: bool, why: str) -> None:
        if not ok:
            raise StepFailed(why)

    def run_step(self, title: str, fn) -> None:
        print(f"\n--- {title} ---", flush=True)
        try:
            fn()
        except StepFailed as exc:
            self.r.check(f"{title}: rest of step skipped", False, str(exc))
        except Exception as exc:  # noqa: BLE001 - report and carry on with the next step
            self.r.check(f"{title}: unexpected error", False, f"{type(exc).__name__}: {exc}")

    # -- 1. Admin and Club

    def step_admin_and_club(self) -> None:
        r = self.r
        result = create_admin(email=ADMIN_EMAIL, name="E2E Admin", password=ADMIN_PASSWORD)
        r.check("create_admin made or kept the Platform Admin", bool(result.student_id))
        admin = Actor(self.api, ADMIN_EMAIL, {})
        self.s["admin"] = admin
        resp = admin.login(ADMIN_PASSWORD)
        self.need(r.status("Admin logs in", resp, 200), "admin login failed")
        r.check("Admin is_platform_admin", resp.json().get("is_platform_admin") is True)

        president = self.actor("president", "CSE", 3, "A")
        coordinator = self.actor("coordinator", "ECE", 3, "B")
        other_president = self.actor("otherpres", "MECH", 2, "A")
        other_coord = self.actor("othercoord", "MECH", 2, "B")
        for a in (president, coordinator, other_president, other_coord):
            self.need(a.signup().status_code == 201, f"signup failed for {a.email}")

        depts = ["CSE", "ECE", "MECH", "CIVIL"]
        sections = ["A", "B", "C"]
        eligible: list[Actor] = []
        for i in range(44):
            eligible.append(self.actor(f"stud{i:02d}", depts[i % 4], 2 + (i % 2), sections[i % 3]))
        ineligible = [
            self.actor(f"inel{i}", depts[i % 4], 1 if i % 2 == 0 else 4, sections[i % 3])
            for i in range(4)
        ]
        failures = [a.email for a in eligible + ineligible if a.signup().status_code != 201]
        r.check(f"{len(eligible) + len(ineligible)} Students sign up", not failures, str(failures))
        self.need(not failures, "student signups failed")
        self.s.update(
            president=president,
            coordinator=coordinator,
            eligible=eligible,
            ineligible=ineligible,
        )

        resp = admin.post(
            "/admin/clubs", json={"name": f"E2E Club {self.suffix}", "slug": f"e2e-{self.suffix}"}
        )
        self.need(r.status("Admin creates a Club", resp, 201), "club create failed")
        club = resp.json()
        self.s["club"] = club
        resp = admin.post(
            "/admin/clubs",
            json={"name": f"Other Club {self.suffix}", "slug": f"e2e-other-{self.suffix}"},
        )
        self.need(r.status("Admin creates a second Club", resp, 201), "club 2 create failed")
        other_club = resp.json()

        resp = admin.put(f"/admin/clubs/{club['id']}/president", json={"student_id": president.id})
        r.status("Admin assigns the President", resp, 200)
        r.check("Assigned role is president", resp.json().get("role") == "president", resp.text)
        resp = admin.put(
            f"/admin/clubs/{other_club['id']}/president", json={"student_id": other_president.id}
        )
        r.status("Admin assigns the other Club's President", resp, 200)

        resp = president.post(
            f"/clubs/{club['id']}/coordinators", json={"student_id": coordinator.id}
        )
        r.status("President adds the Coordinator", resp, 201)
        resp = other_president.post(
            f"/clubs/{other_club['id']}/coordinators", json={"student_id": other_coord.id}
        )
        r.status("Other President adds their Coordinator", resp, 201)

        resp = coordinator.get(f"/clubs/{club['id']}/roles")
        roles = {x["role"] for x in resp.json()} if resp.status_code == 200 else set()
        r.check(
            "Coordinator sees the Club's roles",
            resp.status_code == 200 and roles == {"president", "coordinator"},
            describe(resp),
        )
        resp = other_coord.get(f"/clubs/{club['id']}/roles")
        r.status("Coordinator of another Club gets 403 on roles", resp, 403)
        resp = other_coord.post(
            f"/clubs/{club['id']}/events",
            json={
                "name": "x",
                "starts_at": (datetime.now(UTC) + timedelta(days=5)).isoformat(),
                "ends_at": (datetime.now(UTC) + timedelta(days=5, hours=2)).isoformat(),
            },
        )
        r.status("Coordinator of another Club gets 403 creating an Event", resp, 403)
        self.s["other_coord"] = other_coord

    # -- 2. Event

    def fields(self) -> list[dict[str, Any]]:
        return [
            {"key": "team_name", "label": "Team name", "type": "short_text", "required": True},
            {
                "key": "tshirt",
                "label": "T-shirt size",
                "type": "single_choice",
                "required": True,
                "options": ["S", "M", "L"],
            },
            {
                "key": "interests",
                "label": "Interests",
                "type": "multi_choice",
                "options": ["AI", "Web", "Robotics"],
            },
            {"key": "veg", "label": "Vegetarian meal?", "type": "yes_no", "required": True},
            {"key": "phone", "label": "Phone", "type": "phone"},
            {"key": "notes", "label": "Internal notes", "type": "long_text", "hidden": True},
        ]

    def step_event(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        club = self.s["club"]
        start = datetime.now(UTC) + timedelta(days=7)
        resp = president.post(
            f"/clubs/{club['id']}/events",
            json={
                "name": f"E2E Hackathon {self.suffix}",
                "description": "End-to-end test Event",
                "starts_at": start.isoformat(),
                "ends_at": (start + timedelta(hours=6)).isoformat(),
                "venue": "Main Hall",
                "approval_required": False,
                "seat_limit": 40,
                "registration_deadline": (datetime.now(UTC) + timedelta(days=3)).isoformat(),
            },
        )
        self.need(r.status("President creates the Event", resp, 201), "event create failed")
        event = resp.json()
        self.s["event"] = event
        r.check("Event is open", event.get("is_open") is True, str(event))

        resp = president.put(f"/events/{event['id']}/form", json={"fields": self.fields()})
        self.need(r.status("PUT Registration Form", resp, 200), "form PUT failed")
        form = resp.json()
        r.check(
            "Form keeps all 6 fields, profile fields listed",
            len(form["fields"]) == 6 and "year" in form.get("profile_fields", []),
            str(form),
        )

        resp = president.put(
            f"/events/{event['id']}/eligibility",
            json={"rules": [{"field": "year", "operator": "in", "values": ["2", "3"]}]},
        )
        r.status("PUT Eligibility Rules year in [2, 3]", resp, 200)

        anon = httpx.Client(base_url=self.api, timeout=30)
        resp = anon.get(f"/events/{event['id']}/public")
        anon.close()
        if r.status("Anonymous public page", resp, 200):
            page = resp.json()
            r.check(
                "Anonymous: am_i_eligible and my_registration_status are null",
                page.get("am_i_eligible") is None and page.get("my_registration_status") is None,
                str(page),
            )
            keys = [f["key"] for f in page["form"]["fields"]]
            r.check("Public form hides the hidden field", "notes" not in keys, str(keys))
            r.check("Public page shows 40 seats left", page.get("seats_left") == 40, str(page))

        inel: Actor = self.s["ineligible"][0]
        resp = inel.get(f"/events/{event['id']}/public")
        if r.status("Ineligible Student public page", resp, 200):
            page = resp.json()
            r.check(
                "Ineligible: am_i_eligible false with failed_rules",
                page.get("am_i_eligible") is False and bool(page.get("failed_rules")),
                str(page),
            )
        el: Actor = self.s["eligible"][0]
        resp = el.get(f"/events/{event['id']}/public")
        if r.status("Eligible Student public page", resp, 200):
            page = resp.json()
            r.check(
                "Eligible: am_i_eligible true, failed_rules empty",
                page.get("am_i_eligible") is True and page.get("failed_rules") == [],
                str(page),
            )

    # -- 3. Registrations

    def answers(self, i: int) -> dict[str, Any]:
        return {
            "team_name": "=1+1" if i == 0 else f"Team {i}",
            "tshirt": ["S", "M", "L"][i % 3],
            "interests": [["AI"], ["Web", "Robotics"], []][i % 3],
            "veg": i % 2 == 0,
            "phone": f"+91 98765 {i:05d}",
        }

    def step_registrations(self) -> None:
        r = self.r
        event = self.s["event"]
        eligible: list[Actor] = self.s["eligible"]
        contenders = eligible[:41]
        results: dict[int, httpx.Response] = {}
        barrier = threading.Barrier(len(contenders))

        def go(i: int, a: Actor) -> None:
            with contextlib.suppress(threading.BrokenBarrierError):
                barrier.wait(timeout=30)
            results[i] = a.post(
                f"/events/{event['id']}/registrations", json={"answers": self.answers(i)}
            )

        threads = [threading.Thread(target=go, args=(i, a)) for i, a in enumerate(contenders)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        codes = Counter(resp.status_code for resp in results.values())
        confirmed = [i for i, resp in results.items() if resp.status_code == 201]
        full = [i for i, resp in results.items() if resp.status_code == 409]
        r.check(
            "41 concurrent registrations: exactly 40 confirmed, 1 gets 409",
            len(confirmed) == 40 and len(full) == 1,
            f"status counts {dict(codes)}; "
            + "; ".join(describe(results[i]) for i in results if results[i].status_code >= 500)[
                :400
            ],
        )
        statuses = Counter(results[i].json().get("status") for i in confirmed)
        r.check(
            "All 40 are status confirmed", statuses == Counter({"confirmed": 40}), str(statuses)
        )
        if full:
            code = results[full[0]].json().get("error", {}).get("code")
            r.check(
                "The turned-away Student got a 'full' error",
                code in ("registration_closed", "event_full"),
                describe(results[full[0]]),
            )
        self.need(len(confirmed) >= 2 and len(full) == 1, "seat race result unusable")
        reg_ids = {i: results[i].json()["id"] for i in confirmed}
        self.s["reg_ids"] = reg_ids  # index into eligible -> registration id
        loser = full[0]

        resp = self.s["president"].get(f"/events/{event['id']}")
        r.check(
            "Event shows confirmed_count 40 and closed_reason full",
            resp.status_code == 200
            and resp.json().get("confirmed_count") == 40
            and resp.json().get("closed_reason") == "full",
            describe(resp),
        )

        inel: Actor = self.s["ineligible"][0]
        resp = inel.post(f"/events/{event['id']}/registrations", json={"answers": self.answers(1)})
        r.status("Ineligible Student gets 403", resp, 403)
        r.check(
            "Ineligible error code is not_eligible",
            resp.json().get("error", {}).get("code") == "not_eligible",
            describe(resp),
        )

        # One confirmed Student cancels; the turned-away one takes the seat.
        cancel_i = next(i for i in confirmed if i != 0)
        resp = eligible[cancel_i].delete(f"/events/{event['id']}/registrations/me")
        r.status("A confirmed Student cancels", resp, 204)
        resp = self.s["president"].get(f"/events/{event['id']}")
        r.check(
            "Cancel frees the seat (confirmed_count 39)",
            resp.json().get("confirmed_count") == 39,
            describe(resp),
        )
        # A seat is free now, so answer validation is reached (it runs after the open check).
        bad = {**self.answers(loser), "tshirt": "XXL"}
        resp = eligible[loser].post(f"/events/{event['id']}/registrations", json={"answers": bad})
        r.status("Bad answers (option not listed) give 422", resp, 422)
        bad2 = {"tshirt": "M", "veg": "maybe"}
        resp = eligible[loser].post(f"/events/{event['id']}/registrations", json={"answers": bad2})
        r.status("Bad answers (missing required, bad yes_no) give 422", resp, 422)
        r.check(
            "Bad-answer error code is invalid_answers",
            resp.json().get("error", {}).get("code") == "invalid_answers",
            describe(resp),
        )
        resp = eligible[loser].post(
            f"/events/{event['id']}/registrations", json={"answers": self.answers(loser)}
        )
        r.status("The turned-away Student now registers", resp, 201)
        if resp.status_code == 201:
            reg_ids[loser] = resp.json()["id"]
        del reg_ids[cancel_i]
        self.s["cancelled_i"] = cancel_i

        # A cancelled Student re-registers and gets the same Registration back.
        again_i = next(i for i in reg_ids if i not in (0, loser))
        old_id = reg_ids[again_i]
        resp = eligible[again_i].delete(f"/events/{event['id']}/registrations/me")
        r.status("Another Student cancels", resp, 204)
        resp = eligible[again_i].post(
            f"/events/{event['id']}/registrations", json={"answers": self.answers(again_i)}
        )
        r.status("The cancelled Student re-registers", resp, 201)
        r.check(
            "Re-registration returns the same Registration id",
            resp.status_code == 201 and resp.json().get("id") == old_id,
            f"old {old_id}, got {describe(resp)}",
        )
        r.check(
            "Re-registration is confirmed",
            resp.status_code == 201 and resp.json().get("status") == "confirmed",
            describe(resp),
        )
        resp = eligible[again_i].get(f"/events/{event['id']}/public")
        r.check(
            "Public page shows my_registration_status confirmed",
            resp.status_code == 200 and resp.json().get("my_registration_status") == "confirmed",
            describe(resp),
        )
        resp = eligible[again_i].post(
            f"/events/{event['id']}/registrations", json={"answers": self.answers(again_i)}
        )
        r.status("Registering twice gives 409", resp, 409)

        resp = self.s["president"].get(
            f"/events/{event['id']}/registrations", params={"limit": 1000}
        )
        self.need(r.status("President lists Registrations", resp, 200), "list failed")
        regs = resp.json()
        by_status = Counter(x["status"] for x in regs)
        r.check(
            "Registrations: 40 confirmed, 1 cancelled",
            by_status.get("confirmed") == 40 and by_status.get("cancelled") == 1,
            str(dict(by_status)),
        )
        formula = [x for x in regs if x["answers"].get("team_name") == "=1+1"]
        r.check("One Registration has the =1+1 answer", len(formula) == 1, str(len(formula)))
        self.s["all_regs"] = regs

    # -- 4. Approval

    def step_approval(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        coordinator: Actor = self.s["coordinator"]
        club = self.s["club"]
        start = datetime.now(UTC) + timedelta(days=10)
        resp = president.post(
            f"/clubs/{club['id']}/events",
            json={
                "name": f"E2E Workshop {self.suffix}",
                "starts_at": start.isoformat(),
                "ends_at": (start + timedelta(hours=2)).isoformat(),
                "approval_required": True,
            },
        )
        self.need(r.status("Create an Event with approval on", resp, 201), "event 2 failed")
        ev = resp.json()
        a, b = self.s["eligible"][42], self.s["eligible"][43]
        ra = a.post(f"/events/{ev['id']}/registrations", json={"answers": {}})
        rb = b.post(f"/events/{ev['id']}/registrations", json={"answers": {}})
        r.check(
            "Both Registrations are pending",
            ra.status_code == 201
            and rb.status_code == 201
            and ra.json()["status"] == "pending"
            and rb.json()["status"] == "pending",
            f"{describe(ra)} | {describe(rb)}",
        )
        self.need(ra.status_code == 201 and rb.status_code == 201, "approval regs failed")
        resp = coordinator.post(f"/registrations/{ra.json()['id']}/approve")
        r.check(
            "Coordinator approves: confirmed",
            resp.status_code == 200 and resp.json().get("status") == "confirmed",
            describe(resp),
        )
        resp = president.post(f"/registrations/{rb.json()['id']}/reject")
        r.check(
            "President rejects: rejected",
            resp.status_code == 200 and resp.json().get("status") == "rejected",
            describe(resp),
        )
        resp = b.post(f"/events/{ev['id']}/registrations", json={"answers": {}})
        r.status("Rejected Student can't re-apply (409)", resp, 409)
        resp = president.post(f"/registrations/{rb.json()['id']}/approve")
        r.status("Approving a rejected Registration gives 409", resp, 409)

    # -- 5. Form edit after answers

    def step_form_edit(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        event = self.s["event"]
        base = self.fields()
        optional = {"key": "dietary", "label": "Dietary needs", "type": "short_text"}
        resp = president.put(f"/events/{event['id']}/form", json={"fields": [*base, optional]})
        r.status("Adding an optional field after answers succeeds", resp, 200)
        base = [*base, optional]
        required = {"key": "github", "label": "GitHub", "type": "short_text", "required": True}
        resp = president.put(f"/events/{event['id']}/form", json={"fields": [*base, required]})
        r.status("Adding a required field after answers gives 409", resp, 409)
        without = [f for f in base if f["key"] != "interests"]
        resp = president.put(f"/events/{event['id']}/form", json={"fields": without})
        r.status("Deleting an answered field gives 409", resp, 409)
        hidden = [{**f, "hidden": True} if f["key"] == "interests" else f for f in base]
        resp = president.put(f"/events/{event['id']}/form", json={"fields": hidden})
        r.status("Hiding the answered field succeeds", resp, 200)
        if resp.status_code == 200:
            flags = {f["key"]: f["hidden"] for f in resp.json()["fields"]}
            r.check("The field is now hidden", flags.get("interests") is True, str(flags))

    # -- 6. Poster

    def step_poster(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        event = self.s["event"]
        resp = president.post(
            f"/events/{event['id']}/posters",
            files={"file": ("poster.png", make_poster_png(), "image/png")},
        )
        self.need(r.status("Upload a 2400x3600 PNG Poster", resp, 201), "poster upload failed")
        poster = resp.json()
        r.check(
            "Poster size is 2400x3600",
            poster["width_px"] == 2400 and poster["height_px"] == 3600,
            str(poster),
        )
        pid = poster["id"]

        resp = president.get(f"/posters/{pid}/preview")
        if r.status("GET the Poster preview", resp, 200):
            h = resp.headers
            scale = float(h.get("X-Preview-Scale", "0"))
            r.check(
                "Preview headers: scale and original size",
                h.get("content-type", "").startswith("image/png")
                and h.get("X-Original-Width") == "2400"
                and h.get("X-Original-Height") == "3600"
                and 0 < scale <= 1
                and abs(scale - 1600 / 3600) < 0.01,
                str(dict(h)),
            )

        resp = president.put(
            f"/posters/{pid}/qr-area", json={"x": 100, "y": 100, "w": 200, "h": 200}
        )
        r.status("Too-small QR Area gives 400", resp, 400)
        r.check(
            "Error code is qr_area_too_small",
            resp.json().get("error", {}).get("code") == "qr_area_too_small",
            describe(resp),
        )
        area = {"x": 1750, "y": 2950, "w": 500, "h": 500}
        resp = president.put(f"/posters/{pid}/qr-area", json=area)
        self.need(r.status("Set a valid QR Area", resp, 200), "qr area failed")

        resp = president.post(f"/posters/{pid}/render")
        self.need(r.status("Request the render", resp, 202), "render request failed")

        r.check(
            "Render is queued",
            resp.json().get("render_status") in ("queued", "done"),
            describe(resp),
        )

        # There is no GET for one Poster, so poll the Final Poster: 409 `not_rendered` until
        # the worker is done.
        def fetch() -> httpx.Response:
            return president.get(f"/posters/{pid}/final", params={"format": "png"})

        final = poll(fetch, lambda x: x.status_code != 409, 180, 1.0)
        r.check("The worker renders the Final Poster", final.status_code == 200, describe(final))

        resp = president.get(f"/posters/{pid}/final", params={"format": "png"})
        if r.status("Download the Final Poster PNG", resp, 200):
            r.check("Final PNG is a PNG", resp.content[:8] == b"\x89PNG\r\n\x1a\n")
            want = f"{self.frontend}/events/{event['id']}"
            got = decode_qr(resp.content, area)
            r.check(
                "Final Poster QR decodes to the Event page URL",
                got == want,
                f"got {got!r}, want {want!r}",
            )
        resp = president.get(f"/posters/{pid}/final", params={"format": "pdf"})
        if r.status("Download the Final Poster PDF", resp, 200):
            r.check("Final PDF starts with %PDF", resp.content[:4] == b"%PDF")

        resp = president.get(f"/events/{event['id']}/qr", params={"format": "png"})
        r.check(
            "Event QR as PNG",
            resp.status_code == 200
            and resp.headers.get("content-type", "").startswith("image/png")
            and resp.content[:4] == b"\x89PNG",
            describe(resp) if resp.status_code != 200 else resp.headers.get("content-type", ""),
        )
        if resp.status_code == 200:
            r.check(
                "Event QR PNG decodes to the Event page URL",
                decode_qr(resp.content, {"x": 0, "y": 0, "w": 1, "h": 1})
                == f"{self.frontend}/events/{event['id']}",
            )
        resp = president.get(f"/events/{event['id']}/qr", params={"format": "svg"})
        r.check(
            "Event QR as SVG",
            resp.status_code == 200
            and resp.headers.get("content-type", "").startswith("image/svg+xml")
            and b"<svg" in resp.content,
            describe(resp) if resp.status_code != 200 else resp.headers.get("content-type", ""),
        )

    # -- 7. Export

    def export(self, who: Actor, fmt: str, flt: str) -> tuple[httpx.Response, list[list[Any]]]:
        event = self.s["event"]
        resp = who.get(f"/events/{event['id']}/export", params={"format": fmt, "filter": flt})
        if resp.status_code != 200:
            return resp, []
        rows = csv_rows(resp.content) if fmt == "csv" else xlsx_rows(resp.content)
        return resp, rows

    def step_export(self) -> None:
        r = self.r
        coordinator: Actor = self.s["coordinator"]
        regs = self.s["all_regs"]
        expected = {
            "all": len(regs),
            "confirmed": sum(1 for x in regs if x["status"] == "confirmed"),
            "attendees": sum(1 for x in regs if x["status"] == "confirmed"),
        }
        for fmt in ("csv", "xlsx"):
            for flt in ("all", "confirmed", "attendees"):
                resp, rows = self.export(coordinator, fmt, flt)
                if not r.status(f"Export {fmt} filter={flt}", resp, 200):
                    continue
                r.check(
                    f"Export {fmt} {flt}: {expected[flt]} data rows",
                    len(rows) - 1 == expected[flt],
                    f"got {len(rows) - 1} rows",
                )
                header = [str(h) for h in rows[0]] if rows else []
                r.check(
                    f"Export {fmt} {flt}: hidden field column present",
                    any("Internal notes" in h and "(hidden)" in h for h in header),
                    str(header),
                )
                cells = [c for row in rows[1:] for c in row]
                r.check(
                    f"Export {fmt} {flt}: =1+1 is escaped as '=1+1",
                    "'=1+1" in cells and "=1+1" not in cells,
                    "escaped found" if "'=1+1" in cells else "escaped value missing",
                )
                if fmt == "xlsx" and flt == "all":
                    r.check(
                        "Export xlsx content type",
                        "spreadsheetml" in resp.headers.get("content-type", ""),
                        resp.headers.get("content-type", ""),
                    )
        resp, _ = self.export(self.s["other_coord"], "csv", "all")
        r.status("Coordinator of another Club can't export (403)", resp, 403)

    # -- 8. Attendance

    def step_attendance(self) -> None:
        r = self.r
        coordinator: Actor = self.s["coordinator"]
        event = self.s["event"]
        resp = coordinator.get(f"/events/{event['id']}/attendance")
        self.need(r.status("GET attendance", resp, 200), "attendance GET failed")
        entries = resp.json()
        r.check(
            "Attendance: 40 entries, everyone present",
            len(entries) == 40 and all(e["present"] for e in entries),
            f"{len(entries)} entries, {sum(e['present'] for e in entries)} present",
        )
        reg_ids: dict[int, str] = self.s["reg_ids"]
        # Never mark absent the =1+1 Student (0) or the ones used later.
        absent_i = [i for i in sorted(reg_ids) if i != 0][-3:]
        absent = [reg_ids[i] for i in absent_i]
        resp = coordinator.put(
            f"/events/{event['id']}/attendance", json={"absent_registration_ids": absent}
        )
        self.need(r.status("Mark 3 absent", resp, 200), "attendance PUT failed")
        present = sum(1 for e in resp.json() if e["present"])
        r.check("Attendance now shows 37 present", present == 37, f"{present} present")
        resp, rows = self.export(coordinator, "csv", "attendees")
        r.check(
            "attendees export drops by 3 (37 rows)",
            resp.status_code == 200 and len(rows) - 1 == 37,
            describe(resp) if resp.status_code != 200 else f"{len(rows) - 1} rows",
        )
        self.s["absent_i"] = absent_i
        self.s["attendee_i"] = [i for i in reg_ids if i not in absent_i]

    # -- 9. Certificates

    def step_certificates(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        event = self.s["event"]
        tags = ["name", "event_name", "event_date", "club_name", "form.team_name", "certificate_id"]
        resp = president.post(
            f"/events/{event['id']}/certificate-templates",
            files={"file": ("template.pdf", make_template_pdf(tags), "application/pdf")},
        )
        self.need(r.status("Upload the Certificate Template", resp, 201), "template upload failed")
        tpl = resp.json()
        found = {p["tag"] for p in tpl["placeholders"]}
        r.check("All Placeholders detected", found == set(tags), f"found {sorted(found)}")
        r.check("Template has a text layer", tpl.get("has_text_layer") is True, str(tpl))

        resp = president.post(
            f"/events/{event['id']}/certificate-templates",
            files={
                "file": ("bad.pdf", make_template_pdf(["nmae", "event_name"]), "application/pdf")
            },
        )
        r.status("Template with {{nmae}} gives 400", resp, 400)
        msg = resp.json().get("error", {}).get("message", "")
        r.check("The 400 suggests {{name}}", "{{name}}" in msg, describe(resp))

        resp = president.get(f"/certificate-templates/{tpl['id']}/preview")
        r.check(
            "Template preview is a PDF",
            resp.status_code == 200 and resp.content[:4] == b"%PDF",
            describe(resp) if resp.status_code != 200 else resp.headers.get("content-type", ""),
        )

        subject = "Your certificate for {{event_name}}"
        body = (
            "Hi {{name}},\n\nThanks for joining {{event_name}} by {{club_name}} on "
            "{{event_date}} with team {{form.team_name}}."
        )
        resp = president.post(
            f"/events/{event['id']}/certificate-batches",
            json={"template_id": tpl["id"], "message_subject": subject, "message_body": body},
        )
        self.need(r.status("Create the Certificate Batch", resp, 202), "batch create failed")
        batch = resp.json()
        r.check("Batch total equals 37 Attendees", batch.get("total") == 37, str(batch))
        self.s["batch"] = batch

        state = poll(
            lambda: president.get(f"/certificate-batches/{batch['id']}").json(),
            lambda s: s.get("pending", 1) == 0,
            300,
            2.0,
        )
        r.check(
            "Batch finishes: generated 37, failed 0",
            state.get("pending") == 0 and state.get("generated") == 37 and state.get("failed") == 0,
            str(state),
        )
        r.check("Batch status done", state.get("status") == "done", str(state))
        resp = president.post(f"/certificate-batches/{batch['id']}/retry")
        r.status("Retry with nothing failed gives 409", resp, 409)
        resp = president.post(
            f"/events/{event['id']}/certificate-batches",
            json={"template_id": tpl["id"], "message_subject": subject, "message_body": body},
        )
        r.status("A second batch for the same Attendees gives 409", resp, 409)

    # -- 10. Notifications

    def cert_emails(self, recipients: set[str]) -> list[email.message.EmailMessage]:
        subject = f"Your certificate for {self.s['event']['name']}"
        return [
            m
            for m in read_outbox(self.outbox)
            if str(m["To"]).lower() in recipients and str(m["Subject"]) == subject
        ]

    def step_notifications(self) -> None:
        r = self.r
        eligible: list[Actor] = self.s["eligible"]
        recipients = {eligible[i].email.lower() for i in self.s["attendee_i"]}
        mails = poll(lambda: self.cert_emails(recipients), lambda m: len(m) >= 37, 120, 2.0)
        per = Counter(str(m["To"]).lower() for m in mails)
        r.check(
            "Outbox has exactly one email per Certificate (37)",
            len(mails) == 37 and set(per) == recipients and max(per.values(), default=0) == 1,
            f"{len(mails)} emails, {len(per)} recipients, "
            f"max per person {max(per.values(), default=0)}",
        )
        if not mails:
            return
        problems = []
        for m in mails:
            text = body_text(m)
            to = str(m["To"]).lower()
            actor = next(a for a in eligible if a.email.lower() == to)
            if "{{" in text or "{{" in str(m["Subject"]):
                problems.append(f"{to}: unfilled tag")
            if actor.profile["name"] not in text:
                problems.append(f"{to}: name missing")
            if "/me/certificates" not in text:
                problems.append(f"{to}: no My Certificates link")
            if f"{self.frontend}/verify/" not in text:
                problems.append(f"{to}: no verify link")
            if has_attachment(m):
                problems.append(f"{to}: has an attachment")
        r.check(
            "Every email has tags filled, both links and no attachment",
            not problems,
            "; ".join(problems[:5]),
        )
        sample = body_text(mails[0])
        r.check("Email body includes the team answer", "team" in sample.lower(), sample[:300])

    # -- 11. My Certificates

    def my_certs(self, a: Actor) -> list[dict[str, Any]]:
        resp = a.get("/me/certificates")
        return resp.json() if resp.status_code == 200 else []

    def event_cert(self, a: Actor) -> dict[str, Any]:
        name = self.s["event"]["name"]
        found = [c for c in self.my_certs(a) if c["event_name"] == name]
        if not found:
            raise StepFailed(f"{a.email} has no Certificate for the Event (step 9 failed?)")
        return found[0]

    def step_my_certificates(self) -> None:
        r = self.r
        eligible: list[Actor] = self.s["eligible"]
        attendee_i = self.s["attendee_i"]
        owner = eligible[attendee_i[1]]
        other = eligible[attendee_i[2]]
        resp = owner.get("/me/certificates")
        self.need(r.status("Student lists My Certificates", resp, 200), "list failed")
        certs = [c for c in resp.json() if c["event_name"] == self.s["event"]["name"]]
        r.check("Exactly one Certificate for the Event", len(certs) == 1, str(resp.json()))
        self.need(len(certs) == 1, "no certificate to check")
        cert = certs[0]
        r.check("Certificate is valid", cert["status"] == "valid", str(cert))

        resp = owner.get(f"/me/certificates/{cert['id']}/download")
        self.need(r.status("Owner downloads the PDF", resp, 200), "download failed")
        r.check("Download starts with %PDF", resp.content[:4] == b"%PDF")
        from pypdf import PdfReader

        text = " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(resp.content)).pages)
        r.check("PDF text contains the Student's name", owner.profile["name"] in text, text[:300])
        r.check("PDF text contains the Certificate ID", cert["public_id"] in text, text[:300])
        r.check(
            "PDF text has no {{name}} left", "{{name}}" not in text and "{{" not in text, text[:300]
        )

        resp = other.get(f"/me/certificates/{cert['id']}/download")
        r.status("Another Student downloading it gets 404", resp, 404)

        anon = httpx.Client(base_url=self.api, timeout=30)
        resp = anon.get(f"/verify/{cert['public_id'].lower()}")
        if r.status("Verify by lowercased ID", resp, 200):
            v = resp.json()
            r.check(
                "Verify shows name, event, club, date, status",
                v["name"] == owner.profile["name"]
                and v["event_name"] == self.s["event"]["name"]
                and v["club_name"] == self.s["club"]["name"]
                and v["status"] == "valid"
                and bool(v.get("event_date")),
                str(v),
            )
            leaks = set(v) - VERIFY_FIELDS
            text_v = resp.text
            r.check(
                "Verify has no roll_no, email or answers",
                not leaks
                and owner.profile["roll_no"] not in text_v
                and owner.email not in text_v
                and "answers" not in v,
                f"extra keys {leaks}; body {text_v[:300]}",
            )
        resp = anon.get("/verify/ORG-0000-0000-0000")
        r.status("Unknown Certificate ID gives 404", resp, 404)
        resp = anon.get("/verify/not-a-real-id")
        r.status("Malformed Certificate ID gives 404", resp, 404)
        anon.close()

    # -- 12. Late Attendee

    def all_public_ids(self) -> dict[int, list[str]]:
        eligible: list[Actor] = self.s["eligible"]
        name = self.s["event"]["name"]
        out = {}
        for i in self.s["reg_ids"]:
            out[i] = sorted(
                c["public_id"] for c in self.my_certs(eligible[i]) if c["event_name"] == name
            )
        return out

    def step_late_attendee(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        event = self.s["event"]
        reg_ids: dict[int, str] = self.s["reg_ids"]
        absent_i: list[int] = self.s["absent_i"]
        if "batch" not in self.s:
            raise StepFailed("no Certificate Batch (step 9 failed)")
        before = self.all_public_ids()
        late_i = absent_i[0]
        still_absent = [reg_ids[i] for i in absent_i[1:]]
        resp = president.put(
            f"/events/{event['id']}/attendance", json={"absent_registration_ids": still_absent}
        )
        r.status("Mark one absent Student present again", resp, 200)
        batch = self.s["batch"]
        resp = president.post(
            f"/certificate-batches/{batch['id']}/attendees",
            json={"registration_ids": [reg_ids[late_i]]},
        )
        self.need(r.status("POST the late Attendee to the batch", resp, 202), "late add failed")
        r.check("One item queued", resp.json().get("queued") == 1, describe(resp))
        state = poll(
            lambda: president.get(f"/certificate-batches/{batch['id']}").json(),
            lambda s: s.get("pending", 1) == 0,
            120,
            1.0,
        )
        r.check(
            "Batch now has 38 generated, 0 failed",
            state.get("generated") == 38 and state.get("failed") == 0,
            str(state),
        )
        after = self.all_public_ids()
        new = {i: ids for i, ids in after.items() if ids != before.get(i)}
        r.check(
            "Exactly one new Certificate, for the late Attendee",
            list(new) == [late_i] and len(after[late_i]) == 1 and before[late_i] == [],
            f"changed: {new}",
        )
        r.check(
            "Every other Certificate ID is unchanged",
            all(after[i] == before[i] for i in before if i != late_i),
        )
        resp = president.post(
            f"/certificate-batches/{batch['id']}/attendees",
            json={"registration_ids": [reg_ids[late_i]]},
        )
        r.check(
            "Posting the same late Attendee again creates nothing",
            resp.status_code in (200, 202) and resp.json().get("queued") == 0,
            describe(resp),
        )
        resp = president.post(
            f"/certificate-batches/{batch['id']}/attendees",
            json={"registration_ids": [reg_ids[absent_i[1]]]},
        )
        r.status("Posting a still-absent Student gives 400", resp, 400)
        self.s["attendee_i"] = [*self.s["attendee_i"], late_i]

    # -- 13. Reissue and Revoke

    def step_reissue_revoke(self) -> None:
        r = self.r
        president: Actor = self.s["president"]
        eligible: list[Actor] = self.s["eligible"]
        attendee_i = self.s["attendee_i"]
        name = self.s["event"]["name"]
        anon = httpx.Client(base_url=self.api, timeout=30)

        a = eligible[attendee_i[3]]
        old = self.event_cert(a)
        resp = president.post(f"/certificates/{old['id']}/reissue", json={"reason": "Name typo"})
        self.need(r.status("Reissue a Certificate", resp, 200), "reissue failed")
        mine = [c for c in self.my_certs(a) if c["event_name"] == name]
        newer = [c for c in mine if c["public_id"] != old["public_id"]]
        r.check("A new Certificate ID appears", len(newer) == 1, str(mine))
        new_pid = newer[0]["public_id"] if newer else ""
        resp = anon.get(f"/verify/{old['public_id']}")
        v = resp.json() if resp.status_code == 200 else {}
        r.check(
            "Old ID verifies as superseded, with superseded_by the new ID",
            v.get("status") == "superseded" and v.get("superseded_by") == new_pid,
            describe(resp),
        )
        resp = anon.get(f"/verify/{new_pid}")
        r.check(
            "New ID verifies as valid",
            resp.status_code == 200 and resp.json().get("status") == "valid",
            describe(resp),
        )
        resp = president.post(f"/certificates/{old['id']}/reissue", json={"reason": "again"})
        r.status("Reissuing a superseded Certificate gives 409", resp, 409)

        b = eligible[attendee_i[4]]
        cert = self.event_cert(b)
        resp = president.post(
            f"/certificates/{cert['id']}/revoke", json={"reason": "Did not attend"}
        )
        r.status("Revoke a Certificate", resp, 200)
        resp = anon.get(f"/verify/{cert['public_id']}")
        v = resp.json() if resp.status_code == 200 else {}
        r.check(
            "Revoked ID verifies as revoked with its reason",
            v.get("status") == "revoked" and v.get("revoked_reason") == "Did not attend",
            describe(resp),
        )
        resp = president.post(f"/certificates/{cert['id']}/revoke", json={"reason": "again"})
        r.status("Revoking it again gives 409", resp, 409)
        resp = b.get(f"/me/certificates/{cert['id']}/download")
        r.check(
            "Owner can still download a revoked Certificate, flagged",
            resp.status_code == 200 and resp.headers.get("X-Certificate-Status") == "revoked",
            describe(resp) if resp.status_code != 200 else str(dict(resp.headers)),
        )
        resp = self.s["other_coord"].post(
            f"/certificates/{newer[0]['id'] if newer else old['id']}/revoke", json={"reason": "x"}
        )
        r.check(
            "Coordinator of another Club can't revoke (403/404)",
            resp.status_code in (403, 404),
            describe(resp),
        )
        anon.close()

    # -- 14. Password reset

    def reset_token(self, to: str, seen: set[str]) -> str | None:
        for m in read_outbox(self.outbox):
            if str(m["To"]).lower() != to.lower() or str(m["Message-ID"]) in seen:
                continue
            text = body_text(m)
            if "reset-password?token=" in text:
                return text.split("reset-password?token=", 1)[1].split()[0].strip().rstrip(".")
        return None

    def step_password_reset(self) -> None:
        r = self.r
        eligible: list[Actor] = self.s["eligible"]
        target = eligible[self.s["attendee_i"][5]]
        seen = {
            str(m["Message-ID"])
            for m in read_outbox(self.outbox)
            if str(m["To"]).lower() == target.email.lower()
        }
        anon = httpx.Client(base_url=self.api, timeout=30)
        resp = anon.post("/auth/password-reset", json={"email": target.email})
        r.status("Request a password reset", resp, 202)
        resp = anon.post("/auth/password-reset", json={"email": f"nobody-{self.suffix}@{DOMAIN}"})
        r.status("Reset for an unknown email also gives 202", resp, 202)
        token = poll(lambda: self.reset_token(target.email, seen), lambda t: bool(t), 20, 0.5)
        self.need(r.check("Reset email with a token is in the outbox", bool(token)), "no token")
        resp = anon.post(
            "/auth/password-reset/confirm", json={"token": token, "new_password": NEW_PASSWORD}
        )
        r.status("Confirm the new password", resp, 204)
        fresh = Actor(self.api, target.email, target.profile)
        r.status("Log in with the new password", fresh.login(NEW_PASSWORD), 200)
        r.status(
            "Old password no longer works", Actor(self.api, target.email, {}).login(PASSWORD), 401
        )
        resp = anon.post(
            "/auth/password-reset/confirm", json={"token": token, "new_password": "Another1234"}
        )
        r.status("Reusing the token gives 400", resp, 400)
        r.status("The old session is revoked (old cookie on /me gives 401)", target.get("/me"), 401)
        r.status("The new session works on /me", fresh.get("/me"), 200)
        fresh.close()
        anon.close()

    # -- 15. Removal

    def step_removal(self) -> None:
        r = self.r
        admin: Actor = self.s["admin"]
        eligible: list[Actor] = self.s["eligible"]
        victim = eligible[self.s["attendee_i"][6]]
        name = self.s["event"]["name"]
        certs = [c for c in self.my_certs(victim) if c["event_name"] == name]
        resp = admin.delete(f"/admin/students/{victim.id}")
        r.status("Admin removes a Student who holds a Certificate", resp, 204)
        r.status("Removed Student's old session gives 401", victim.get("/me"), 401)
        r.status(
            "Removed Student can't log in", Actor(self.api, victim.email, {}).login(PASSWORD), 401
        )
        if certs:
            anon = httpx.Client(base_url=self.api, timeout=30)
            resp = anon.get(f"/verify/{certs[0]['public_id']}")
            r.check(
                "Their Certificate still verifies as valid",
                resp.status_code == 200 and resp.json().get("status") == "valid",
                describe(resp),
            )
            anon.close()
        else:
            r.check(
                "Their Certificate still verifies as valid",
                False,
                "they hold no Certificate, so this can't be checked (step 9 failed)",
            )
        resp = admin.get("/admin/students", params={"roll_no": victim.profile["roll_no"]})
        r.check(
            "Admin search shows them as removed",
            resp.status_code == 200 and [s["status"] for s in resp.json()] == ["removed"],
            describe(resp),
        )

    # -- run

    def run(self) -> int:
        print(f"ORGniser e2e run {self.suffix} against {self.api}")
        print(f"Outbox: {self.outbox}")
        health = httpx.get(self.api + "/health", timeout=10)
        if health.status_code != 200:
            print(f"Server not healthy: {describe(health)}")
            return 2
        steps = [
            ("1. Admin and Club", self.step_admin_and_club),
            ("2. Event", self.step_event),
            ("3. Registrations", self.step_registrations),
            ("4. Approval", self.step_approval),
            ("5. Form edit after answers", self.step_form_edit),
            ("6. Poster", self.step_poster),
            ("7. Export", self.step_export),
            ("8. Attendance", self.step_attendance),
            ("9. Certificates", self.step_certificates),
            ("10. Notifications", self.step_notifications),
            ("11. My Certificates", self.step_my_certificates),
            ("12. Late Attendee", self.step_late_attendee),
            ("13. Reissue and Revoke", self.step_reissue_revoke),
            ("14. Password reset", self.step_password_reset),
            ("15. Removal", self.step_removal),
        ]
        started = time.monotonic()
        for title, fn in steps:
            self.run_step(title, fn)
        for a in self.actors:
            a.close()
        print(f"\nRun {self.suffix} took {time.monotonic() - started:.0f} s")
        return self.r.summary()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run the whole product end to end over HTTP.")
    parser.add_argument("--base-url", default="http://localhost:8002")
    parser.add_argument(
        "--outbox",
        required=True,
        type=Path,
        help="The server's ORGNISER_EMAIL_OUTBOX_DIR, where emails are saved as .eml files.",
    )
    parser.add_argument(
        "--frontend-base-url",
        default="http://localhost:3000",
        help="The server's ORGNISER_FRONTEND_BASE_URL; the Event QR must point here.",
    )
    args = parser.parse_args(argv)
    return E2E(args.base_url, args.outbox, args.frontend_base_url).run()


if __name__ == "__main__":
    raise SystemExit(main())
