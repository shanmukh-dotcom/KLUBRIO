"""Pure tests for app.exports.format (R23-R24; engineering rule 6)."""

import csv
import io
import time
from datetime import UTC, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app.events.schemas import FormField
from app.exports.format import (
    build_columns,
    build_rows,
    escape_cell,
    export_filename,
    to_csv,
    to_xlsx,
)

DANGEROUS = ["=1+1", "+91 98765", "-2", "@SUM(A1)", "\tx", "\rx", '=HYPERLINK("http://x","y")']


def make_fields():
    return [
        {
            "key": "tshirt",
            "label": "T-shirt",
            "type": "single_choice",
            "position": 2,
            "hidden": False,
        },
        {
            "key": "langs",
            "label": "Languages",
            "type": "multi_choice",
            "position": 1,
            "hidden": False,
        },
        {
            "key": "old_q",
            "label": "Old question",
            "type": "short_text",
            "position": 3,
            "hidden": True,
        },
        {"key": "veg", "label": "Vegetarian", "type": "yes_no", "position": 4, "hidden": False},
    ]


def make_reg(**overrides):
    reg = {
        "profile_snapshot": {
            "name": "Ánanyā Rāo",
            "roll_no": "21CS001",
            "department": "CSE",
            "year": 3,
            "section": "A",
            "captured_at": "2026-10-01T00:00:00+00:00",
        },
        "email": "ananya@uni.edu",
        "answers": {"tshirt": "M", "langs": ["py", "go"], "old_q": "legacy", "veg": True},
        "status": "confirmed",
        "submitted_at": datetime(
            2026, 10, 1, 15, 30, tzinfo=timezone(timedelta(hours=5, minutes=30))
        ),
        "attended": None,
        "certificate_id": None,
        "certificate_status": None,
    }
    reg.update(overrides)
    return reg


# --------------------------------------------------------------------------- escape_cell


@pytest.mark.parametrize("value", DANGEROUS)
def test_escape_dangerous_prefixes(value):
    assert escape_cell(value) == "'" + value


@pytest.mark.parametrize("value", ["Asha", "", "1=1", "a-b"])
def test_escape_leaves_safe_strings(value):
    assert escape_cell(value) == value


@pytest.mark.parametrize("value", [-5, 3.2, True, datetime(2026, 1, 1)])
def test_escape_leaves_non_strings(value):
    assert escape_cell(value) is value


def test_escape_none_becomes_empty():
    assert escape_cell(None) == ""


# --------------------------------------------------------------------------- columns


def test_column_order_and_hidden_included():
    cols = build_columns(make_fields())
    assert [c.key for c in cols] == [
        "name",
        "roll_no",
        "department",
        "year",
        "section",
        "email",
        "langs",
        "tshirt",
        "old_q",
        "veg",
        "status",
        "registered_at",
        "attended",
        "certificate_id",
        "certificate_status",
    ]
    headers = {c.key: c.header for c in cols}
    assert headers["old_q"] == "Old question (hidden)"
    assert headers["tshirt"] == "T-shirt"


def test_columns_accept_pydantic_fields_without_position():
    fields = [
        FormField(key="b", label="B", type="short_text"),
        FormField(key="a", label="A", type="short_text", hidden=True),
    ]
    keys = [c.key for c in build_columns(fields) if c.source == "answer"]
    assert keys == ["b", "a"]


# --------------------------------------------------------------------------- rows


def test_row_formatting():
    cols = build_columns(make_fields())
    row = dict(zip([c.key for c in cols], next(build_rows([make_reg()], cols)), strict=True))
    assert row["name"] == "Ánanyā Rāo"
    assert row["year"] == 3
    assert row["langs"] == "py; go"
    assert row["veg"] == "Yes"
    assert row["old_q"] == "legacy"
    assert row["registered_at"] == "2026-10-01T10:00:00+00:00"
    assert row["attended"] == ""
    assert row["certificate_id"] == ""


def test_row_attended_and_certificate():
    cols = build_columns(make_fields())
    reg = make_reg(
        attended=False,
        certificate_id="ORG-ABC-123",
        certificate_status="valid",
        answers={"veg": False},
    )
    row = dict(zip([c.key for c in cols], next(build_rows([reg], cols)), strict=True))
    assert row["attended"] == "No"
    assert row["veg"] == "No"
    assert row["tshirt"] == ""
    assert row["certificate_id"] == "ORG-ABC-123"
    assert row["certificate_status"] == "valid"


def test_naive_datetime_treated_as_utc():
    cols = build_columns([])
    reg = make_reg(submitted_at=datetime(2026, 10, 1, 10, 0))
    row = dict(zip([c.key for c in cols], next(build_rows([reg], cols)), strict=True))
    assert row["registered_at"] == "2026-10-01T10:00:00+00:00"


def test_rows_escape_every_cell():
    cols = build_columns(make_fields())
    reg = make_reg(
        email="=evil@x",
        answers={"tshirt": "+1", "langs": ["-a", "b"], "old_q": "@x"},
        profile_snapshot={"name": "=cmd|'/c calc'!A0", "roll_no": "-1"},
    )
    row = dict(zip([c.key for c in cols], next(build_rows([reg], cols)), strict=True))
    for key in ("email", "tshirt", "langs", "old_q", "name", "roll_no"):
        assert row[key].startswith("'"), key


# --------------------------------------------------------------------------- CSV


def test_csv_has_bom_and_headers():
    cols = build_columns(make_fields())
    data = to_csv(cols, build_rows([make_reg()], cols))
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text)))
    assert parsed[0][0] == "Name"
    assert "Old question (hidden)" in parsed[0]
    assert parsed[1][0] == "Ánanyā Rāo"


@pytest.mark.parametrize("value", DANGEROUS)
def test_csv_escapes_dangerous_prefixes(value):
    cols = build_columns([{"key": "q", "label": "Q", "type": "short_text", "position": 0}])
    reg = make_reg(answers={"q": value})
    text = to_csv(cols, build_rows([reg], cols)).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text)))
    q_index = parsed[0].index("Q")
    assert parsed[1][q_index] == "'" + value


def test_csv_escapes_dangerous_header_and_raw_rows():
    cols = build_columns([{"key": "q", "label": "=Q", "type": "short_text", "position": 0}])
    text = to_csv(cols, [["=raw"] * len(cols)]).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text)))
    assert "'=Q" in parsed[0]
    assert all(cell == "'=raw" for cell in parsed[1])


# --------------------------------------------------------------------------- XLSX


def read_xlsx(data):
    wb = load_workbook(io.BytesIO(data))
    return wb.active


@pytest.mark.parametrize("value", DANGEROUS)
def test_xlsx_escapes_dangerous_prefixes(value):
    cols = build_columns([{"key": "q", "label": "Q", "type": "short_text", "position": 0}])
    reg = make_reg(answers={"q": value})
    ws = read_xlsx(to_xlsx(cols, build_rows([reg], cols)))
    headers = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=headers.index("Q") + 1)
    assert cell.data_type == "s"
    assert str(cell.value).startswith("'")


def test_xlsx_header_bold_frozen_widths():
    cols = build_columns(make_fields())
    ws = read_xlsx(to_xlsx(cols, build_rows([make_reg()], cols)))
    assert ws.freeze_panes == "A2"
    assert all(c.font.bold for c in ws[1])
    assert [c.value for c in ws[1]][:6] == [
        "Name",
        "Roll No",
        "Department",
        "Year",
        "Section",
        "Email",
    ]
    assert ws.column_dimensions["A"].width == cols[0].width
    assert ws.cell(row=2, column=1).value == "Ánanyā Rāo"
    assert ws.cell(row=2, column=4).value == 3


def test_xlsx_5000_rows_is_fast():
    fields = make_fields()
    cols = build_columns(fields)
    regs = [make_reg() for _ in range(5000)]
    start = time.perf_counter()
    data = to_xlsx(cols, build_rows(regs, cols))
    elapsed = time.perf_counter() - start
    assert elapsed < 5, elapsed
    ws = read_xlsx(data)
    assert ws.max_row == 5001


# --------------------------------------------------------------------------- filename


def test_export_filename():
    now = datetime(2026, 10, 25, 23, 0, tzinfo=timezone(timedelta(hours=-5)))
    assert export_filename("Tech Fest 2026!", "xlsx", now=now) == (
        "tech-fest-2026_registrations_2026-10-26.xlsx"
    )
    assert export_filename("../../etc/passwd", "CSV", now=now).startswith("etc-passwd_")
    assert export_filename("नमस्ते", "csv", now=now) == "event_registrations_2026-10-26.csv"


def test_export_filename_default_date_is_utc_today():
    assert datetime.now(UTC).date().isoformat() in export_filename("x", "csv")


def test_export_filename_rejects_unknown_format():
    with pytest.raises(ValueError):
        export_filename("x", "pdf")
